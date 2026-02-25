"""
투자경고종목 연도별 분리 수집 스크립트

2020~2026년(또는 지정 연도)의 투자경고종목을 연도별로 수집하여
output/ 디렉토리에 CSV 및 Excel 파일로 저장합니다.

사용법:
  uv run python collect_yearly_warnings.py           # 2020~2026 전체
  uv run python collect_yearly_warnings.py --year 2023         # 특정 연도만
  uv run python collect_yearly_warnings.py --start 2022 --end 2024  # 범위 지정
"""

import argparse
import os
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from investment_hub.core.ports.storage_port import StoragePort
from investment_hub.infrastructure.adapters.local_storage_adapter import LocalStorageAdapter
from investment_hub.infrastructure.collectors.daily_price_collector import collect_daily_prices_batch
from investment_hub.infrastructure.scrapers.krx_warning_scraper import fetch_investment_warning_stocks

# ─────────────────────────────────────────────────────────────────────────────
# 설정
# ─────────────────────────────────────────────────────────────────────────────
OUTPUT_DIR = "output"
MAX_WARNING_DAYS = 60  # 60일 이하 경고 종목만 수집
TRADING_DAYS_AFTER_RELEASE = 3  # 해제일 이후 추가 수집일

# 기본 저장소로 LocalStorageAdapter 사용
_storage: StoragePort = LocalStorageAdapter()


def set_storage(storage_adapter: StoragePort):
    global _storage
    _storage = storage_adapter


def get_storage() -> StoragePort:
    return _storage


# ─────────────────────────────────────────────────────────────────────────────
# 연도별 날짜 범위 (현재 연도는 오늘 날짜 자동 반영)
# ─────────────────────────────────────────────────────────────────────────────
def _build_year_ranges() -> dict:
    today = datetime.now().strftime("%Y-%m-%d")
    current_year = datetime.now().year
    ranges = {}
    for y in range(2020, current_year):
        ranges[y] = (f"{y}-01-01", f"{y}-12-31")
    ranges[current_year] = (f"{current_year}-01-01", today)
    return ranges


YEAR_RANGES = _build_year_ranges()


# ─────────────────────────────────────────────────────────────────────────────
# 핵심 함수
# ─────────────────────────────────────────────────────────────────────────────


def collect_year(year: int, include_active: bool = False) -> bool:
    """
    특정 연도의 투자경고종목을 수집하고 CSV + Excel로 저장합니다.

    Args:
        year: 수집 연도
        include_active: True이면 아직 해제되지 않은 진행 중 종목도 포함

    Returns True if successful, False otherwise.
    """
    if year not in YEAR_RANGES:
        print(f"  [오류] {year}년은 지원하지 않는 연도입니다.")
        return False

    start_date, end_date = YEAR_RANGES[year]
    print(f"\n{'=' * 60}")
    print(f"  {year}년 투자경고종목 수집 시작")
    print(f"  기간: {start_date} ~ {end_date}")
    print(f"  진행 중 종목 포함: {'예' if include_active else '아니오'}")
    print(f"{'=' * 60}")

    # ── Step 1: 투자경고 지정 목록 수집 ──────────────────────────────────────
    print("\n[1/4] KRX KIND에서 투자경고 목록 수집 중...")
    stock_list = fetch_investment_warning_stocks(start_date, end_date)
    print(f"  → 총 {len(stock_list)}건 수집")

    if not stock_list:
        print("  데이터 없음. 건너뜀.")
        return False

    # ── Step 2: 필터링 ────────────────────────────────────────────────────────
    print(f"\n[2/4] 필터링 (코넥스 제외, {MAX_WARNING_DAYS}일 이하)...")
    start_dt = pd.to_datetime(start_date)
    end_dt = pd.to_datetime(end_date)

    filtered = []
    for s in stock_list:
        if s.designation_date is None:
            continue
        if not (start_dt <= s.designation_date <= end_dt):
            continue
        if s.market == "코넥스":
            continue
        # 해제된 종목: 경고일 60일 이하만
        if s.release_date is not None:
            if (s.release_date - s.designation_date).days > MAX_WARNING_DAYS:
                continue
        else:
            # 진행 중 종목: include_active 플래그가 있어야 포함
            if not include_active:
                continue
        filtered.append(s)

    print(f"  → 필터 후 {len(filtered)}건")

    if not filtered:
        print("  필터 후 데이터 없음. 건너뜀.")
        return False

    # ── Step 3: 일별 가격 수집 ───────────────────────────────────────────────
    print("\n[3/4] 일별 가격 데이터 수집 (market_daily 캐시 활용)...")
    daily_prices_by_code = collect_daily_prices_batch(
        filtered,
        trading_days_after_release=TRADING_DAYS_AFTER_RELEASE,
    )

    # ── Step 4: 저장 ─────────────────────────────────────────────────────────
    print("\n[4/4] 결과 저장 (StorageAdapter 사용)...")
    storage = get_storage()
    storage.ensure_directory(OUTPUT_DIR)

    _save_csv(year, filtered, daily_prices_by_code, storage)
    _save_excel(year, filtered, daily_prices_by_code, storage)

    print(f"\n  [완료] {year}년 수집 완료")
    return True


def _build_rows(filtered: list, daily_prices_by_code: dict) -> tuple[list, int]:
    """
    Excel/CSV 공통 rows_data 생성.
    Returns (rows_data, max_days)
    """
    max_days = 0
    rows_data = []

    for stock_info in filtered:
        code = stock_info.code
        daily_prices = daily_prices_by_code.get(code, [])
        if not daily_prices:
            continue

        daily_prices_sorted = sorted(daily_prices, key=lambda x: x.date)
        designation_date = stock_info.designation_date
        release_date = stock_info.release_date

        # 해제일 인덱스 탐색
        release_trading_day_orig = None
        for idx, dp in enumerate(daily_prices_sorted):
            if release_date and dp.date.date() == release_date.date():
                release_trading_day_orig = idx
                break
        if release_date and release_trading_day_orig is None:
            for idx, dp in enumerate(daily_prices_sorted):
                if dp.date.date() >= release_date.date():
                    release_trading_day_orig = idx
                    break

        # valid_prices: 지정일부터 해제일+3일
        valid_prices = []
        for idx, dp in enumerate(daily_prices_sorted):
            if designation_date and dp.date.date() < designation_date.date():
                continue
            if release_trading_day_orig is not None and idx > release_trading_day_orig + TRADING_DAYS_AFTER_RELEASE:
                break
            valid_prices.append(dp)

        prices_by_trading_day = {}
        new_release_trading_day = None
        for i, dp in enumerate(valid_prices):
            prices_by_trading_day[i] = {
                "date": dp.date,
                "close": dp.close,
                "change_rate": dp.change_rate,
            }
            # 해제일 탐색:
            # 1) 해제일 당일 정상 거래(종가 > 0)가 있으면 그날을 해제 거래일로 사용
            # 2) 해제일 당일 거래정지(데이터 없거나 종가=0)이면 그 이후 첫 정상 거래일을 사용
            if release_date:
                if dp.date.date() == release_date.date() and dp.close > 0:
                    new_release_trading_day = i  # 해제일 당일 정상 거래
                elif new_release_trading_day is None and dp.date.date() > release_date.date() and dp.close > 0:
                    # 해제일이 거래정지 → 해제 이후 첫 정상 거래일
                    new_release_trading_day = i

        if prices_by_trading_day:
            max_days = max(max_days, len(prices_by_trading_day) - 1)

        rows_data.append(
            {
                "name": stock_info.name,
                "code": code,
                "market": stock_info.market,
                "designation_date": designation_date.strftime("%Y-%m-%d") if designation_date else "",
                "release_date": release_date.strftime("%Y-%m-%d") if release_date else "진행중",
                "warning_days": (release_date - designation_date).days if release_date else None,
                "prices_by_trading_day": prices_by_trading_day,
                "release_trading_day": new_release_trading_day,
            }
        )

    return rows_data, max_days


def _calc_returns(row_data: dict) -> tuple:
    """
    해제전/해제직후 등락률 계산.
    해제일(release_day)은 이미 정상거래가 가능한 날이므로:
    - pre_return (해제전): 지정일(0) 대비 해제 전날(release_day - 1)
    - post_return (해제직후): 해제 전날 대비 해제일(release_day)
    """
    pre_return = None
    post_return = None
    release_day = row_data["release_trading_day"]
    ptd = row_data["prices_by_trading_day"]

    # 해제일이 지정일 이후여야 함 (최소 D+1)
    if release_day is not None and release_day > 0:
        prev_day = release_day - 1

        # 1. 해제전 등락률 (지정일 종가 -> 해제 전날 종가)
        if 0 in ptd and prev_day in ptd:
            d0_price = ptd[0]["close"]
            prev_price = ptd[prev_day]["close"]
            if d0_price > 0:
                pre_return = round((prev_price / d0_price - 1) * 100, 2)

        # 2. 해제직후 등락률 (해제 전날 종가 -> 해제일 종가)
        if prev_day in ptd and release_day in ptd:
            prev_price = ptd[prev_day]["close"]
            rel_price = ptd[release_day]["close"]
            if prev_price > 0:
                post_return = round((rel_price / prev_price - 1) * 100, 2)

    return pre_return, post_return


def _save_csv(year: int, filtered: list, daily_prices_by_code: dict, storage: StoragePort):
    """연도별 상세 CSV 저장 (long format)."""
    records = []
    for stock_info in filtered:
        code = stock_info.code
        prices = daily_prices_by_code.get(code, [])
        for dp in sorted(prices, key=lambda x: x.date):
            records.append(
                {
                    "year": year,
                    "code": code,
                    "name": stock_info.name,
                    "market": stock_info.market,
                    "designation_date": stock_info.designation_date.strftime("%Y-%m-%d")
                    if stock_info.designation_date
                    else "",
                    "release_date": stock_info.release_date.strftime("%Y-%m-%d") if stock_info.release_date else "",
                    "date": dp.date.strftime("%Y-%m-%d"),
                    "close": dp.close,
                    "change_rate": dp.change_rate,
                }
            )

    if not records:
        print("  CSV: 저장할 데이터 없음")
        return

    df = pd.DataFrame(records)
    csv_path = os.path.join(OUTPUT_DIR, f"투자경고종목분석({year}년).csv")
    storage.save_dataframe_csv(df, csv_path)


def _save_excel(year: int, filtered: list, daily_prices_by_code: dict, storage: StoragePort):
    """연도별 Wide format Excel 저장."""
    rows_data, max_days = _build_rows(filtered, daily_prices_by_code)

    if not rows_data:
        print("  Excel: 저장할 데이터 없음")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = f"투자경고_{year}"

    # ── 헤더 ──────────────────────────────────────────────────────────────────
    base_headers = [
        "종목명",
        "종목코드",
        "시장",
        "지정일",
        "해제일",
        "경고일수",
        "해제전등락률(%)",
        "해제직후등락률(%)",
    ]
    day_headers = [f"D+{d}" for d in range(max_days + 1)]
    header = base_headers + day_headers
    ws.append(header)

    # 헤더 스타일
    header_fill = PatternFill(start_color="2F4F8F", end_color="2F4F8F", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # ── 데이터 행 ────────────────────────────────────────────────────────────
    green_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
    red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
    BASE_COL = len(base_headers)  # D+0 이 시작되는 컬럼 오프셋

    for row_data in rows_data:
        pre_return, post_return = _calc_returns(row_data)

        row = [
            row_data["name"],
            row_data["code"],
            row_data["market"],
            row_data["designation_date"],
            row_data["release_date"],
            row_data["warning_days"],
            pre_return,
            post_return,
        ]

        ptd = row_data["prices_by_trading_day"]
        for d in range(max_days + 1):
            row.append(ptd[d]["close"] if d in ptd else None)

        ws.append(row)

        # 셀 배경 강조
        current_row = ws.max_row
        for d in range(max_days + 1):
            if d not in ptd:
                continue
            col_idx = BASE_COL + d + 1
            cell = ws.cell(row=current_row, column=col_idx)
            change_rate = ptd[d]["change_rate"]
            if change_rate >= 29.9:
                cell.fill = red_fill
            elif d == row_data["release_trading_day"]:
                cell.fill = green_fill

    # ── 열 너비 자동 조정 ────────────────────────────────────────────────────
    for col_idx, col_cells in enumerate(ws.columns, 1):
        max_len = max((len(str(c.value or "")) for c in col_cells), default=0)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 20)

    # 틀 고정 (헤더 + 기본 컬럼)
    ws.freeze_panes = ws.cell(row=2, column=BASE_COL + 1)

    xlsx_path = os.path.join(OUTPUT_DIR, f"투자경고종목분석({year}년).xlsx")
    storage.save_workbook(wb, xlsx_path)


# ─────────────────────────────────────────────────────────────────────────────
# 진입점
# ─────────────────────────────────────────────────────────────────────────────


def parse_args():
    parser = argparse.ArgumentParser(
        description="투자경고종목 연도별 수집",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
예시:
  uv run python collect_yearly_warnings.py               # 2020~현재 전체
  uv run python collect_yearly_warnings.py --year 2025  # 2025년만
  uv run python collect_yearly_warnings.py --start 2023 --end 2025
  uv run python collect_yearly_warnings.py --include-active  # 진행 중 종목 포함
""",
    )
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--year", type=int, help="특정 연도만 수집 (예: 2023)")
    group.add_argument("--start", type=int, default=2020, help="시작 연도 (기본: 2020)")
    parser.add_argument("--end", type=int, default=datetime.now().year, help="종료 연도 (기본: 현재 연도)")
    parser.add_argument("--include-active", action="store_true", help="해제일 없는 진행 중 종목도 포함")
    return parser.parse_args()


def main():
    args = parse_args()
    include_active = args.include_active

    if args.year:
        years = [args.year]
    else:
        years = list(range(args.start, args.end + 1))

    # 지원하지 않는 연도 필터 (YEAR_RANGES에 없는 연도)
    valid_years = [y for y in years if y in YEAR_RANGES]
    skipped = set(years) - set(valid_years)
    if skipped:
        print(f"[경고] 수집 불가 연도 제외: {sorted(skipped)}")

    print(f"\n{'=' * 60}")
    print("  투자경고종목 연도별 수집")
    print(f"  대상 연도: {valid_years}")
    print(f"  진행 중 종목 포함: {'예' if include_active else '아니오'}")
    print(f"{'=' * 60}")

    results = {}
    for year in valid_years:
        ok = collect_year(year, include_active=include_active)
        results[year] = "[완료]" if ok else "[실패/데이터없음]"

    print(f"\n{'=' * 60}")
    print("  최종 결과 요약")
    print(f"{'=' * 60}")
    for year, status in results.items():
        print(f"  {year}년: {status}")
    print(f"\n  출력 디렉토리: {os.path.abspath(OUTPUT_DIR)}/")
    print(f"{'=' * 60}\n")


if __name__ == "__main__":
    main()
