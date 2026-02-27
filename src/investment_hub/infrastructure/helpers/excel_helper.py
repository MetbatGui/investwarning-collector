import os
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from investment_hub.core.ports.storage_port import StoragePort
from investment_hub.domain.models import DailyPriceData, InvestmentWarningStock
from investment_hub.domain.services import calculate_returns


def restore_stocks_and_prices_from_df(df: pd.DataFrame):
    """CSV DataFrame에서 종목 및 시세 도메인 모델 복원"""
    if df.empty:
        return [], {}

    stocks_dict = {}
    for _, row in df.drop_duplicates(subset=["code", "designation_date"]).iterrows():
        stocks_dict[row["code"]] = InvestmentWarningStock(
            code=str(row["code"]),
            name=str(row["name"]),
            market=str(row["market"]),
            designation_date=pd.to_datetime(row["designation_date"]),
            release_date=pd.to_datetime(row["release_date"])
            if pd.notna(row["release_date"]) and row["release_date"] != ""
            else None,
        )

    prices_dict = {}
    for code, group in df.groupby("code"):
        prices_dict[str(code)] = [
            DailyPriceData(
                code=str(code),
                name=str(row["name"]),
                date=pd.to_datetime(row["date"]),
                close=float(row["close"]),
                change_rate=float(row["change_rate"]),
            )
            for _, row in group.iterrows()
        ]
    return list(stocks_dict.values()), prices_dict


def _prepare_excel_rows(stocks: list, prices_by_code: dict) -> tuple[list, int]:
    rows, max_days = [], 0
    for s in sorted(stocks, key=lambda x: x.designation_date or datetime.min):
        prices = sorted(prices_by_code.get(s.code, []), key=lambda x: x.date)
        if not prices: continue
        ptd = {i: {"close": p.close, "change_rate": p.change_rate} for i, p in enumerate(prices)}
        max_days = max(max_days, max(ptd.keys()) if ptd else 0)
        rel_idx = next((i for i, p in enumerate(prices) if s.release_date and p.date.date() == s.release_date.date()), None)
        pre, post = calculate_returns(ptd, rel_idx)
        rows.append({"stock": s, "ptd": ptd, "release_idx": rel_idx, "pre_return": pre, "post_return": post,
                     "warning_days": (s.release_date - s.designation_date).days if s.release_date else None})
    return rows, max_days

def _write_excel_headers(ws, max_days: int) -> list[str]:
    base = ["종목명", "종목코드", "시장", "지정일", "해제일", "경고일수", "지정이후 수익률", "해제이후 수익률"]
    ws.append(base + [f"D+{d}" for d in range(max_days + 1)])
    fill, font = PatternFill("solid", "2F4F8F"), Font("FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment = fill, font, Alignment("center")
    return base

def _write_excel_data_rows(ws, rows_data: list, max_days: int, base_cnt: int):
    r_fill, g_fill = PatternFill("solid", "FFCCCC"), PatternFill("solid", "CCFFCC")
    for data in rows_data:
        s = data["stock"]
        row = [s.name, s.code, s.market, s.designation_date.strftime("%Y-%m-%d") if s.designation_date else "",
               s.release_date.strftime("%Y-%m-%d") if s.release_date else "진행중", data["warning_days"],
               data["pre_return"], data["post_return"]]
        row.extend(data["ptd"].get(d, {}).get("close") for d in range(max_days + 1))
        ws.append(row)
        for d in range(max_days + 1):
            if d in data["ptd"]:
                cell = ws.cell(ws.max_row, base_cnt + d + 1)
                if data["ptd"][d]["change_rate"] >= 29.9: cell.fill = r_fill
                if d == data["release_idx"]: cell.fill = g_fill

def _finalize_excel_sheet(ws, base_cnt: int):
    for col in ws.columns:
        lengths = [len(str(c.value or "")) for c in col]
        ws.column_dimensions[col[0].column_letter].width = min(max(lengths) + 2, 30)
    ws.freeze_panes = ws.cell(2, base_cnt + 1)

def save_investment_warning_excel(year: int, stocks: list, prices_by_code: dict, storage: StoragePort, output_dir: str = "output"):
    """투자경고 종목 분석 데이터를 엑셀로 저장 (Rich Format)"""
    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}년 투자경고 분석"

    rows_data, max_days = _prepare_excel_rows(stocks, prices_by_code)
    base_headers = _write_excel_headers(ws, max_days)
    _write_excel_data_rows(ws, rows_data, max_days, len(base_headers))
    _finalize_excel_sheet(ws, len(base_headers))

    storage.save_workbook(wb, os.path.join(output_dir, f"투자경고종목분석({year}년).xlsx"))
